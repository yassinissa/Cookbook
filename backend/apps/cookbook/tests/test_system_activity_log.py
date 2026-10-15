"""
System-wide action log (SystemActivityLog via apps.cookbook.audit) and the
Activity export.

Pins: item-supplement saves are logged with an itemised before → after (one
row per conversion line); a save that changes nothing writes no row; deletes
survive the thing they delete (recipe deletes included); the feed merges these
rows with the recipe logs and filters / scopes / strips money the same way;
users' passwords are never logged; the .xlsx / .csv export carries every
matching row in the caller's scope.
"""
import csv
import io
from decimal import Decimal
from unittest import mock

from django.contrib.admin.models import CHANGE, LogEntry
from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.core.cache import cache
from django.core.management import CommandError, call_command
from openpyxl import load_workbook
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.test import APIClient, APITestCase

from apps.accounts.models import Capability, Role
from apps.cookbook.models import (
    ActivityActionType, Allergen, Branch, DigestSubscription, DishRecipe, DishRecipeActivityLog, Menu, Section,
    SystemActivityLog,
)
from .support import make_units

User = get_user_model()
FEED = '/api/cookbook/activity/'
CONV = '/api/cookbook/item-conversions/'


def fields(row):
    return {f['field']: (f['from'], f['to']) for f in row.changes.get('fields', [])}


class SystemLogTestBase(APITestCase):
    def setUp(self):
        self.units = make_units()
        self.g = self.units['g']
        self.salmiya = Branch.objects.create(name_en='Salmiya', sort_order=1)
        self.jabriya = Branch.objects.create(name_en='Jabriya', sort_order=2)
        self.section = Section.objects.create(name='Salad', avg_monthly_salary=Decimal('285'))
        self.admin = User.objects.create_superuser('boss', password='x')
        self.client = self.api(self.admin)
        patcher = mock.patch('apps.cookbook.publishing.sync_item_conversions',
                             return_value={'updated': [], 'skipped': [], 'problems': []})
        patcher.start()
        self.addCleanup(patcher.stop)

    def api(self, user):
        c = APIClient(HTTP_ACCEPT='application/json')
        c.force_authenticate(user)
        return c

    def _user(self, role_name, branches=(), extra=()):
        u = User.objects.create_user(f'u{User.objects.count()}', password='x')
        p = u.profile
        p.role = Role.objects.get(name=role_name)
        p.scope_overridden = bool(branches)
        p.save()
        p.branches.set(branches)
        p.extra_capabilities.set(Capability.objects.filter(code__in=extra))
        return u

    def _conversion(self, sku='B900', qty='15'):
        return self.client.post(CONV, {
            'item_sku': sku, 'base_unit': str(self.g.id), 'order_cost': '2.000', 'pack_qty': '1000',
            'lines': [{'label': '1 Tbs', 'quantity': qty, 'unit': str(self.g.id)}],
        }, format='json')


class ItemSupplementLogTests(SystemLogTestBase):
    def test_create_update_delete_conversion_are_logged(self):
        self.assertEqual(self._conversion().status_code, 201)
        created = SystemActivityLog.objects.get(action_type=ActivityActionType.CREATED)
        self.assertEqual((created.area, created.subject_type, created.subject_code),
                         ('item', 'item_conversion', 'B900'))
        self.assertEqual(created.changed_by, 'boss')
        self.assertEqual(fields(created)['conversion.line 1 Tbs'], (None, '15 g'))
        self.assertEqual(created.subject_path, '/inventory?q=B900')

        r = self.client.patch(f'{CONV}B900/', {
            'lines': [{'label': '1 Tbs', 'quantity': '18', 'unit': str(self.g.id)},
                      {'label': '1 Cup', 'quantity': '240', 'unit': str(self.g.id)}],
        }, format='json')
        self.assertEqual(r.status_code, 200, r.content)
        updated = SystemActivityLog.objects.get(action_type=ActivityActionType.UPDATED)
        self.assertEqual(fields(updated), {
            'conversion.line 1 Tbs': ('15 g', '18 g'),
            'conversion.line 1 Cup': (None, '240 g'),
        })

        self.assertEqual(self.client.delete(f'{CONV}B900/').status_code, 204)
        deleted = SystemActivityLog.objects.get(action_type=ActivityActionType.DELETED)
        self.assertEqual(deleted.subject_code, 'B900')
        self.assertEqual(fields(deleted)['conversion.line 1 Tbs'], ('18 g', None))

    def test_a_save_that_changes_nothing_writes_no_row(self):
        self._conversion()
        self.client.patch(f'{CONV}B900/', {
            'lines': [{'label': '1 Tbs', 'quantity': '15.000', 'unit': str(self.g.id)}],
        }, format='json')
        self.assertFalse(SystemActivityLog.objects.filter(action_type=ActivityActionType.UPDATED).exists())

    def test_nutrition_and_storage_are_logged(self):
        self.client.post('/api/cookbook/item-nutrition/', {
            'item_sku': 'B900', 'unit_scale': str(self.g.id), 'calories': '4'}, format='json')
        self.client.post('/api/cookbook/item-storage/', {
            'item_sku': 'B900', 'storage_band': 'chilled', 'shelf_life_hours': 48}, format='json')
        self.client.patch('/api/cookbook/item-storage/B900/', {'shelf_life_hours': 24}, format='json')
        types = list(SystemActivityLog.objects.order_by('created_at')
                     .values_list('subject_type', 'action_type'))
        self.assertEqual(types, [('item_nutrition', 'created'), ('item_storage', 'created'),
                                 ('item_storage', 'updated')])
        upd = SystemActivityLog.objects.get(subject_type='item_storage', action_type='updated')
        self.assertEqual(fields(upd), {'storage.shelf_life_hours': ('48', '24')})


class FeedAndScopeTests(SystemLogTestBase):
    def test_feed_merges_system_rows_with_recipe_rows(self):
        dish = DishRecipe.objects.create(name_en='Tabbouleh', recipe_code='T1',
                                         branch_ref=self.salmiya, section=self.section)
        DishRecipeActivityLog.objects.create(recipe=dish, action_type='created', changed_by='boss')
        self._conversion()

        r = self.client.get(FEED)
        self.assertEqual(r.status_code, 200)
        self.assertEqual([e['kind'] for e in r.data['results']], ['item', 'dish'])
        item = r.data['results'][0]
        self.assertEqual(item['recipe_code'], 'B900')
        self.assertEqual(item['subject_type'], 'item_conversion')
        self.assertTrue(item['changes']['fields'])
        self.assertIn('item', {k['value'] for k in r.data['kinds']})

        r = self.client.get(FEED, {'kind': 'item'})
        self.assertEqual(r.data['count'], 1)
        r = self.client.get(FEED, {'q': 'b900'})
        self.assertEqual(r.data['count'], 1)

    def test_recipe_delete_is_logged_and_outlives_the_recipe(self):
        dish = DishRecipe.objects.create(name_en='Fattoush', recipe_code='F1',
                                         branch_ref=self.salmiya, section=self.section)
        self.assertEqual(self.client.delete(f'/api/cookbook/dish-recipes/{dish.id}/').status_code, 204)
        r = self.client.get(FEED, {'kind': 'dish'})
        self.assertEqual([(e['action'], e['recipe_name']) for e in r.data['results']],
                         [('deleted', 'Fattoush')])

    def test_item_rows_need_inventory_view_and_money_is_stripped(self):
        self._conversion()
        staff = self._user('Branch Staff', branches=[self.salmiya], extra=['activity.view'])
        self.assertEqual(self.api(staff).get(FEED).data['count'], 0)

        qa = self._user('QA Manager')           # inventory.view, no costing.view
        self.assertFalse(Role.objects.get(name='QA Manager').capabilities.filter(code='costing.view').exists())
        entry = self.api(qa).get(FEED).data['results'][0]
        names = {f['field'] for f in entry['changes']['fields']}
        self.assertIn('conversion.line 1 Tbs', names)
        self.assertNotIn('conversion.order_cost', names)
        self.assertNotIn('conversion.cost_per_base_unit', names)

    def test_menu_rows_are_branch_scoped(self):
        sal = Menu.objects.create(branch=self.salmiya)
        jab = Menu.objects.create(branch=self.jabriya)
        for m in (sal, jab):
            r = self.client.post('/api/cookbook/menu-periods/', {
                'menu': str(m.id), 'kind': 'seasonal', 'name_en': f'Ramadan {m.branch.name_en}',
                'starts_on': '2026-03-01'}, format='json')
            self.assertEqual(r.status_code, 201, r.content)
        cook = self._user('Restaurant Cook', branches=[self.salmiya], extra=['activity.view', 'menu.view'])
        names = {e['recipe_name'] for e in self.api(cook).get(FEED, {'kind': 'menu'}).data['results']}
        self.assertEqual(names, {'Ramadan Salmiya'})

    def test_user_password_is_never_logged(self):
        r = self.client.post('/api/accounts/users/', {
            'username': 'nadia', 'password': 'hunter2-secret', 'email': 'n@example.com'}, format='json')
        self.assertEqual(r.status_code, 201, r.content)
        uid = r.data['id']
        self.client.patch(f'/api/accounts/users/{uid}/', {'password': 'another-secret'}, format='json')
        self.client.post(f'/api/accounts/users/{uid}/deactivate/')

        rows = SystemActivityLog.objects.filter(subject_type='user').order_by('created_at')
        self.assertEqual([r.action_type for r in rows], ['created', 'updated', 'deactivated'])
        self.assertEqual(fields(rows[0])['user.password'], (None, 'set'))
        self.assertEqual(fields(rows[1]), {'user.password': (None, 'changed')})
        blob = str(list(rows.values_list('changes', 'description')))
        self.assertNotIn('hunter2', blob)
        self.assertNotIn('another-secret', blob)


class ExportTests(SystemLogTestBase):
    def test_xlsx_has_an_activity_sheet_and_a_field_change_sheet(self):
        self._conversion()
        self.client.patch(f'{CONV}B900/', {
            'lines': [{'label': '1 Tbs', 'quantity': '18', 'unit': str(self.g.id)}]}, format='json')

        r = self.client.get(f'{FEED}export/', {'kind': 'item'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('attachment;', r['Content-Disposition'])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Activity', 'Field changes'])
        activity = list(wb['Activity'].iter_rows(values_only=True))
        self.assertEqual(activity[0][:4], ('Date', 'Time', 'Area', 'Action'))
        self.assertEqual(len(activity), 3)                      # header + updated + created
        self.assertEqual(activity[1][3], 'Updated')
        self.assertIn('1 Tbs: 15 g → 18 g', activity[1][10])
        changes = list(wb['Field changes'].iter_rows(values_only=True))
        self.assertIn(('conversion · line 1 Tbs', '15 g', '18 g'),
                      [row[7:10] for row in changes])

    def test_csv_and_validation(self):
        self._conversion()
        r = self.client.get(f'{FEED}export/', {'file': 'csv'})
        self.assertEqual(r.status_code, 200)
        rows = list(csv.reader(io.StringIO(r.content.decode('utf-8-sig'))))
        self.assertEqual(rows[0][0], 'Date')
        self.assertEqual(len(rows), 2)
        self.assertEqual(self.client.get(f'{FEED}export/', {'file': 'pdf'}).status_code, 400)

    def test_export_is_capability_gated(self):
        staff = self._user('Branch Staff', branches=[self.salmiya])
        self.assertEqual(self.api(staff).get(f'{FEED}export/').status_code, 403)


class SignInLogTests(SystemLogTestBase):
    LOGIN = '/api/auth/login/'

    def setUp(self):
        super().setUp()
        cache.clear()
        self.addCleanup(cache.clear)
        self.nadia = User.objects.create_user('nadia', password='right-horse-battery')
        self.anon = APIClient(HTTP_ACCEPT='application/json', REMOTE_ADDR='10.1.2.3')

    def test_sign_in_and_failed_sign_in_are_logged_without_the_password(self):
        r = self.anon.post(self.LOGIN, {'username': 'nadia', 'password': 'right-horse-battery'}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertIn('access', r.data)
        r = self.anon.post(self.LOGIN, {'username': 'nadia', 'password': 'wrong-guess-123'}, format='json')
        self.assertEqual(r.status_code, 401)

        ok = SystemActivityLog.objects.get(action_type=ActivityActionType.SIGNED_IN)
        self.assertEqual((ok.area, ok.changed_by, ok.subject_code), ('auth', 'nadia', 'nadia'))
        self.assertIn('10.1.2.3', ok.description)
        bad = SystemActivityLog.objects.get(action_type=ActivityActionType.SIGN_IN_FAILED)
        self.assertEqual((bad.subject_code, bad.changed_by), ('nadia', ''))
        blob = str(list(SystemActivityLog.objects.values_list('description', 'changes', 'subject_label')))
        self.assertNotIn('right-horse-battery', blob)
        self.assertNotIn('wrong-guess-123', blob)

    def test_sign_out_is_logged(self):
        token = self.anon.post(self.LOGIN, {'username': 'nadia', 'password': 'right-horse-battery'},
                               format='json').data['access']
        c = APIClient(HTTP_ACCEPT='application/json', HTTP_AUTHORIZATION=f'Bearer {token}')
        self.assertEqual(c.post('/api/auth/logout/').status_code, 204)
        row = SystemActivityLog.objects.get(action_type=ActivityActionType.SIGNED_OUT)
        self.assertEqual((row.area, row.changed_by), ('auth', 'nadia'))
        self.assertEqual(APIClient().post('/api/auth/logout/').status_code, 401)

    def test_rate_limited_attempt_is_logged(self):
        with mock.patch.object(ScopedRateThrottle, 'THROTTLE_RATES', {'login': '1/min'}):
            self.anon.post(self.LOGIN, {'username': 'x', 'password': 'y'}, format='json')
            r = self.anon.post(self.LOGIN, {'username': 'x', 'password': 'y'}, format='json')
        self.assertEqual(r.status_code, 429)
        self.assertTrue(SystemActivityLog.objects.filter(
            action_type='sign_in_failed', description__startswith='Blocked').exists())

    def test_sign_ins_are_visible_to_user_admins_only(self):
        self.anon.post(self.LOGIN, {'username': 'nadia', 'password': 'right-horse-battery'}, format='json')
        self.assertEqual(self.client.get(FEED, {'kind': 'auth'}).data['count'], 1)
        qa = self._user('QA Manager')
        self.assertEqual(self.api(qa).get(FEED, {'kind': 'auth'}).data['count'], 0)


class OtherActionLogTests(SystemLogTestBase):
    def test_unsubscribe_link_is_logged(self):
        sub = DigestSubscription.objects.create(user=self.admin, cadence='weekly')
        APIClient().get(f'/api/cookbook/public/digest/unsubscribe/{sub.unsubscribe_token}/')
        row = SystemActivityLog.objects.get(subject_type='digest_subscription')
        self.assertEqual((row.changed_by, row.description), ('boss', 'Unsubscribed via the email link'))

    def test_command_runs_are_logged_with_output(self):
        call_command('sync_capabilities', stdout=io.StringIO())
        row = SystemActivityLog.objects.get(subject_type='command')
        self.assertEqual((row.area, row.action_type, row.subject_code, row.changed_by),
                         ('system', 'command_run', 'sync_capabilities', 'system'))
        self.assertTrue(row.description.startswith('manage.py sync_capabilities'))

    def test_failed_command_is_logged_and_still_raises(self):
        with self.assertRaises(CommandError):
            call_command('import_pos_menu', 'no-such-file.xlsm', '--dry-run', stdout=io.StringIO())
        row = SystemActivityLog.objects.get(action_type=ActivityActionType.COMMAND_FAILED)
        self.assertIn('--dry-run', row.description)
        self.assertIn('File not found', fields(row)['command.error'][1])

    def test_django_admin_edits_appear_in_the_feed(self):
        allergen = Allergen.objects.create(name='Sesame')
        LogEntry.objects.log_action(
            user_id=self.admin.pk, content_type_id=ContentType.objects.get_for_model(Allergen).pk,
            object_id=str(allergen.pk), object_repr='Sesame', action_flag=CHANGE,
            change_message='[{"changed": {"fields": ["Name"]}}]')
        r = self.client.get(FEED, {'kind': 'admin'})
        e = r.data['results'][0]
        self.assertEqual((e['action'], e['recipe_name'], e['changed_by']), ('updated', 'Sesame', 'boss'))
        self.assertIn('Django admin', e['subject_type_display'])
        x = self.client.get(f'{FEED}export/', {'kind': 'admin'})
        self.assertEqual(x.status_code, 200)
